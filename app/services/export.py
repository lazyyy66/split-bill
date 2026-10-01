"""Экспорт группы в Excel: траты (с долями каждого), балансы, как рассчитаться, история переводов.

Суммы пишутся числами (Decimal в основной валюте) с денежным форматом — в Excel их можно складывать
и сортировать. Даты — в часовом поясе группы, без tzinfo (Excel его не поддерживает).
"""

import re
from datetime import datetime
from decimal import Decimal
from io import BytesIO
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.bot.texts import t
from app.db.models import Expense, Group, Settlement
from app.domain.money import CURRENCIES
from app.domain.settlement import minimize_transfers
from app.services.balances import group_balances
from app.services.expenses import list_categories
from app.services.groups import list_members

HEADER_FILL = PatternFill("solid", fgColor="18181C")
HEADER_FONT = Font(bold=True, color="FFFFFF")
TOTAL_FONT = Font(bold=True)
ACCENT_FONT = Font(bold=True, color="C21FD1")  # пурпурный бренда, чуть темнее — читается на белом

SHEETS = {
    "ru": ("Траты", "Балансы", "Как рассчитаться", "Переводы"),
    "en": ("Expenses", "Balances", "Settle up", "Transfers"),
}
COLUMNS = {
    "ru": {
        "date": "Дата",
        "title": "Название",
        "category": "Категория",
        "amount": "Сумма",
        "payer": "Платил(а)",
        "member": "Участник",
        "paid": "Заплатил(а)",
        "share": "Его доля",
        "sent": "Перевёл(а)",
        "received": "Получил(а)",
        "balance": "Баланс",
        "from": "Кто",
        "to": "Кому",
        "details": "Реквизиты",
        "status": "Статус",
        "total": "Итого",
        "left": "вышел(а)",
        "settled": "Все в расчёте",
    },
    "en": {
        "date": "Date",
        "title": "Title",
        "category": "Category",
        "amount": "Amount",
        "payer": "Paid by",
        "member": "Member",
        "paid": "Paid",
        "share": "Share",
        "sent": "Sent",
        "received": "Received",
        "balance": "Balance",
        "from": "From",
        "to": "To",
        "details": "Payment details",
        "status": "Status",
        "total": "Total",
        "left": "left",
        "settled": "Everyone is settled",
    },
}
STATUSES = {
    "ru": {"pending": "ждёт подтверждения", "confirmed": "подтверждён", "rejected": "отклонён"},
    "en": {"pending": "pending", "confirmed": "confirmed", "rejected": "rejected"},
}


def export_filename(group: Group, now: datetime) -> str:
    title = re.sub(r'[\\/:*?"<>|\s]+', "_", group.title).strip("_")[:40] or "group"
    return f"SplitBill_{title}_{now:%Y-%m-%d}.xlsx"


async def build_export(session: AsyncSession, group: Group, *, timezone: str) -> bytes:
    lang = group.language if group.language in SHEETS else "en"
    col = COLUMNS[lang]
    currency = CURRENCIES[group.currency]
    unit = Decimal(10) ** currency.exponent
    money_format = (
        f'#,##0.{"0" * currency.exponent} "{currency.symbol}"' if currency.exponent else f'#,##0 "{currency.symbol}"'
    )
    tz = ZoneInfo(timezone)

    def major(amount: int) -> Decimal:
        return Decimal(amount) / unit

    def local(dt: datetime) -> datetime:
        return dt.astimezone(tz).replace(tzinfo=None)

    members = await list_members(session, group, include_left=True)
    names = {m.user_id: m.user.name for m in members}
    member_ids = [m.user_id for m in members]
    categories = {c.id: c for c in await list_categories(session, group)}
    expenses = list(
        await session.scalars(
            select(Expense)
            .where(Expense.group_id == group.id, Expense.deleted_at.is_(None))
            .order_by(Expense.created_at)
        )
    )
    settlements = list(
        await session.scalars(select(Settlement).where(Settlement.group_id == group.id).order_by(Settlement.created_at))
    )
    balances = await group_balances(session, group)

    def category_name(category_id: int) -> str:
        category = categories.get(category_id)
        if category is None:
            return ""
        label = t(lang, f"cat_{category.code}") if category.code else category.name
        return f"{category.emoji} {label}"

    wb = Workbook()
    expenses_ws = wb.active
    assert expenses_ws is not None
    expenses_ws.title, balances_title, settle_title, transfers_title = SHEETS[lang]

    # --- траты: по колонке на каждого участника ---
    header = [col["date"], col["title"], col["category"], col["amount"], col["payer"], *(names[i] for i in member_ids)]
    expenses_ws.append(header)
    for expense in expenses:
        shares = {s.user_id: s.amount for s in expense.shares}
        expenses_ws.append(
            [
                local(expense.created_at),
                expense.title,
                category_name(expense.category_id),
                major(expense.amount),
                names.get(expense.payer_id, "?"),
                *(major(shares[i]) if i in shares else None for i in member_ids),
            ]
        )
    if expenses:
        last = len(expenses) + 1
        total_row = [col["total"], None, None, f"=SUM(D2:D{last})", None]
        total_row += [
            f"=SUM({get_column_letter(6 + n)}2:{get_column_letter(6 + n)}{last})" for n in range(len(member_ids))
        ]
        expenses_ws.append(total_row)
        for cell in expenses_ws[expenses_ws.max_row]:
            cell.font = TOTAL_FONT
    _format_sheet(expenses_ws, money_columns=range(4, 6 + len(member_ids)), money_format=money_format, date_column=1)

    # --- балансы ---
    paid: dict[int, int] = {}
    owed: dict[int, int] = {}
    for expense in expenses:
        paid[expense.payer_id] = paid.get(expense.payer_id, 0) + expense.amount
        for share in expense.shares:
            owed[share.user_id] = owed.get(share.user_id, 0) + share.amount
    sent: dict[int, int] = {}
    received: dict[int, int] = {}
    for s in settlements:
        if s.status == "confirmed":
            sent[s.from_user_id] = sent.get(s.from_user_id, 0) + s.amount
            received[s.to_user_id] = received.get(s.to_user_id, 0) + s.amount

    balances_ws = wb.create_sheet(balances_title)
    balances_ws.append([col["member"], col["paid"], col["share"], col["sent"], col["received"], col["balance"]])
    for member in sorted(members, key=lambda m: -balances.get(m.user_id, 0)):
        uid = member.user_id
        label = names[uid] + (f" ({col['left']})" if member.left_at else "")
        balances_ws.append(
            [
                label,
                major(paid.get(uid, 0)),
                major(owed.get(uid, 0)),
                major(sent.get(uid, 0)),
                major(received.get(uid, 0)),
                major(balances.get(uid, 0)),
            ]
        )
        if balances.get(uid, 0) > 0:
            balances_ws.cell(balances_ws.max_row, 6).font = ACCENT_FONT
    _format_sheet(balances_ws, money_columns=range(2, 7), money_format=money_format)

    # --- как рассчитаться ---
    settle_ws = wb.create_sheet(settle_title)
    settle_ws.append([col["from"], col["to"], col["amount"], col["details"]])
    transfers = minimize_transfers(balances)
    details = {m.user_id: m.user.payment_details for m in members}
    for tr in transfers:
        settle_ws.append([names[tr.from_user], names[tr.to_user], major(tr.amount), details.get(tr.to_user) or ""])
    if not transfers:
        settle_ws.append([col["settled"]])
    _format_sheet(settle_ws, money_columns=[3], money_format=money_format)

    # --- история переводов ---
    transfers_ws = wb.create_sheet(transfers_title)
    transfers_ws.append([col["date"], col["from"], col["to"], col["amount"], col["status"]])
    for s in settlements:
        transfers_ws.append(
            [
                local(s.created_at),
                names.get(s.from_user_id, "?"),
                names.get(s.to_user_id, "?"),
                major(s.amount),
                STATUSES[lang].get(s.status, s.status),
            ]
        )
    _format_sheet(transfers_ws, money_columns=[4], money_format=money_format, date_column=1)

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _format_sheet(ws: Worksheet, *, money_columns, money_format: str, date_column: int | None = None) -> None:
    """Шапка тёмная и закреплена, автофильтр, деньги и даты в своём формате, ширина колонок по содержимому."""
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"
    if ws.max_row > 1:
        ws.auto_filter.ref = ws.dimensions

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            if cell.column in money_columns:
                cell.number_format = money_format
            elif cell.column == date_column and isinstance(cell.value, datetime):
                cell.number_format = "dd.mm.yyyy hh:mm"

    for column_cells in ws.columns:
        width = max(
            (len(str(c.value)) for c in column_cells if c.value is not None and not str(c.value).startswith("=")),
            default=8,
        )
        ws.column_dimensions[get_column_letter(column_cells[0].column)].width = min(max(width + 3, 10), 45)
