"""Экспорт в Excel: содержимое таблицы и доставка файла в личку."""

from decimal import Decimal
from io import BytesIO

import pytest
from openpyxl import load_workbook

from tests.integration.tg import TgHarness, find_button


@pytest.fixture
async def trip(tg: TgHarness):
    """Роман — ужин 9000 на троих; Даша вернула Роману 3000 и вышла; Арман — такси 1500 на двоих."""
    roman, arman, dasha = tg.person("Роман", "roman"), tg.person("Арман", "arman"), tg.person("Даша", "dasha")
    chat = tg.group("Алматы")
    [welcome] = await tg.add_bot_to_group(chat, by=roman)
    for person in (arman, dasha):
        await person.click(welcome, "Я в деле")
    await roman.send(roman.private_chat, "/setpay Kaspi +7 777 123 45 67")
    await roman.send(chat, "/add 9000 ужин")
    [request] = await dasha.send(chat, "/paid 3000 @roman")
    await roman.click(request, "Получил")
    await tg.leave_group(dasha, chat)  # Даша рассчиталась и вышла — такси делим на двоих
    await arman.send(chat, "/add 1500 такси")
    tg.telegram.reset()
    return chat, roman, arman, dasha


def sheet_rows(data: bytes, title: str) -> list[tuple]:
    wb = load_workbook(BytesIO(data))
    return [tuple(row) for row in wb[title].iter_rows(values_only=True)]


async def test_export_is_sent_privately_with_all_sheets(tg: TgHarness, trip):
    chat, roman, arman, dasha = trip
    [reply] = await roman.send(chat, "/export")
    assert "Отправил таблицу в личку" in reply.text

    [doc] = tg.telegram.documents
    assert doc.chat_id == roman.tg.id  # в личку, не в группу
    assert doc.filename.startswith("SplitBill_Алматы_") and doc.filename.endswith(".xlsx")
    assert "Алматы" in (doc.caption or "")

    wb = load_workbook(BytesIO(doc.data))
    assert wb.sheetnames == ["Траты", "Балансы", "Как рассчитаться", "Переводы"]

    expenses = sheet_rows(doc.data, "Траты")
    assert expenses[0][:5] == ("Дата", "Название", "Категория", "Сумма", "Платил(а)")
    assert set(expenses[0][5:]) == {"Роман", "Арман", "Даша"}
    titles = [row[1] for row in expenses[1:-1]]
    assert titles == ["ужин", "такси"]
    dinner = expenses[1]
    assert dinner[2] == "🍽 Кафе и рестораны"
    assert isinstance(dinner[3], int | float | Decimal) and dinner[3] == 9000  # число, а не строка
    assert expenses[-1][0] == "Итого" and str(expenses[-1][3]).startswith("=SUM(")


async def test_export_balances_and_settle_up(tg: TgHarness, trip):
    chat, roman, arman, dasha = trip
    await roman.send(chat, "/export")
    data = tg.telegram.documents[0].data

    balances = {row[0]: row[1:] for row in sheet_rows(data, "Балансы")[1:]}
    # Роман: заплатил 9000, доли 3000 (ужин) + 750 (такси), получил 3000 от Даши → +2250
    assert balances["Роман"] == (9000, 3750, 0, 3000, 2250)
    # Арман: заплатил 1500, доли 3000 + 750 → −2250
    assert balances["Арман"] == (1500, 3750, 0, 0, -2250)
    # Даша вышла, но в таблице есть: доля 3000, перевела 3000 → 0
    assert balances["Даша (вышел(а))"] == (0, 3000, 3000, 0, 0)

    settle = sheet_rows(data, "Как рассчитаться")
    assert settle[1:] == [("Арман", "Роман", 2250, "Kaspi +7 777 123 45 67")]

    transfers = sheet_rows(data, "Переводы")
    assert transfers[1][1:] == ("Даша", "Роман", 3000, "подтверждён")


async def test_export_needs_private_chat(tg: TgHarness, trip):
    chat, roman, arman, dasha = trip
    [reply] = await arman.send(chat, "/export")  # Арман ни разу не писал боту в личку
    assert "Не могу написать тебе в личку" in reply.text
    assert find_button(reply.reply_markup, "Открыть чат с ботом").url.endswith("?start=export")
    assert tg.telegram.documents == []


async def test_export_in_english(tg: TgHarness, trip):
    chat, roman, *_ = trip
    await roman.send(chat, "/lang en")
    await roman.send(chat, "/export")
    wb = load_workbook(BytesIO(tg.telegram.documents[0].data))
    assert wb.sheetnames == ["Expenses", "Balances", "Settle up", "Transfers"]
